from docx import Document
import sys

sys.stdout.reconfigure(encoding='utf-8')

doc = Document('docs/Al Jamea Grade 6 English Test Set 1 - Final.docx')

content = []
for para in doc.paragraphs:
    if para.text.strip():
        content.append(para.text)

for table in doc.tables:
    for row in table.rows:
        row_data = [cell.text for cell in row.cells]
        content.append(' | '.join(row_data))

for line in content:
    print(line)
